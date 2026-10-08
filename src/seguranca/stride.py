"""Matriz de Ameaças e Riscos (STRIDE) da Rede Alimenta IA: fonte única de dados.

Vocabulário do professor: Ativo -> Ameaça -> Vulnerabilidade -> Impacto -> Controle -> Evidência -> Reteste.
Risco = Probabilidade (1-5) x Impacto (1-5). Faixas: 1-4 Baixo, 5-9 Médio, 10-15 Alto, 16-25 Crítico.
Inerente = sem controles; Residual = com os controles listados (implementados + planejados).

Gerar planilha e tabela:
    python -m src.seguranca.stride
"""

from __future__ import annotations

from dataclasses import asdict, dataclass

from src.data.curadoria_ipvs import RAIZ

DIR_SAIDA = RAIZ / "docs" / "stride"


@dataclass
class Ameaca:
    id: str
    componente: str
    stride: str
    ameaca: str
    vulnerabilidade: str
    ativo: str
    principio: str
    referencia: str
    p: int
    i: int
    preventivo: str
    detectivo: str
    corretivo: str
    status: str
    p_res: int
    i_res: int
    evidencia: str

    @property
    def inerente(self) -> int:
        return self.p * self.i

    @property
    def residual(self) -> int:
        return self.p_res * self.i_res


def nivel(risco: int) -> str:
    if risco >= 16:
        return "Crítico"
    if risco >= 10:
        return "Alto"
    if risco >= 5:
        return "Médio"
    return "Baixo"


AMEACAS: list[Ameaca] = [
    # ---------------- S: Spoofing (falsificação de identidade) ----------------
    Ameaca("S1", "Cadastro de ONG (API)", "S - Spoofing",
           "ONG de fachada se cadastra para desviar alimentos e dinheiro do caixa solidário",
           "Cadastro aberto sem verificação da entidade", "Alimentos, caixa, confiança dos doadores",
           "Integridade / Governança", "—", 4, 5,
           "Dígito verificador do CNPJ (inclusive alfanumérico) + consulta à Receita (situação ATIVA, natureza jurídica sem fins lucrativos) + aprovação manual do Admin (regras 6.3); só ONG aprovada usa o caixa",
           "Log de aprovação/reprovação; alerta de ONG com pedidos muito acima do que recebe (linha de base, regras 6.5/10.2)",
           "Suspensão imediata e revisão dos lotes recebidos",
           "Implementado (v2.0)", 2, 4,
           "src/validacao/documentos.py, receita.py; tests/validacao/test_receita.py; ongs.csv (3 pendentes, 2 reprovadas fora do matching)"),
    Ameaca("S2", "Coleta (transportador)", "S - Spoofing",
           "Pessoa se passa por transportador cadastrado e retira o lote (roubo de carga)",
           "Identidade do transportador não conferida no momento da coleta", "Alimentos, segurança do doador",
           "Integridade", "—", 3, 4,
           "Cadastro verificado do transportador; código de coleta (PIN) exibido ao transportador e conferido pelo doador; endereço exato só após o aceite",
           "Evento COLETA com ator identificado; divergência PIN x transportador gera alerta",
           "Bloqueio do transportador e registro do incidente",
           "Parcial (log + endereço após aceite); PIN planejado CP3", 2, 3,
           "eventos.csv.gz (COLETA com ator); regras §3 e §6.2"),
    Ameaca("S3", "Autenticação (Supabase Auth)", "S - Spoofing",
           "Roubo de credencial de ONG ou do gestor do caixa (phishing) para aceitar lotes ou aprovar gastos",
           "Senha como único fator; sessões longas", "Contas, caixa, dados de lotes",
           "Confidencialidade / Integridade", "MITRE ATT&CK T1566 (Phishing)", 3, 4,
           "MFA obrigatório para ONG, gestor e admin; token de sessão curto (15 min, aula 7); rate limit no login",
           "Alerta de login de local/dispositivo novo; log de autenticação",
           "Revogação de sessões e reset de credenciais",
           "Planejado CP3", 2, 3, "—"),
    Ameaca("S5", "Consulta à Receita (BrasilAPI)", "S - Spoofing",
           "Resposta falsificada ou adulterada (DNS/MITM, serviço comprometido) faz um CNPJ inativo parecer ativo",
           "Confiar na resposta externa sem validar; aprovar por omissão quando o serviço cai", "Verificação de ONG",
           "Integridade", "OWASP LLM05 (saída externa não confiável)", 2, 4,
           "HTTPS sem seguir redirecionamento; resposta validada por esquema; CNPJ da resposta tem de ser o consultado; falha fechada (PENDENTE); aprovação humana final",
           "Log de consultas PENDENTES e respostas fora do esquema",
           "Reconsulta e revisão manual",
           "Implementado (v2.0)", 1, 3,
           "src/validacao/receita.py; tests/validacao/test_receita.py"),
    Ameaca("S4", "Integração com app de entrega", "S - Spoofing",
           "Callback falso confirma uma entrega paga que não aconteceu",
           "Webhook sem autenticação da origem", "Caixa solidário",
           "Integridade", "—", 2, 4,
           "Webhook assinado (HMAC) e validado; a entrega só conta com confirmação da ONG",
           "Conciliação diária: entregas pagas x eventos ENTREGA confirmados",
           "Estorno/contestação junto ao app e bloqueio do fluxo",
           "Planejado CP3", 1, 4, "regras §9 (livro-caixa com comprovante)"),

    # ---------------- T: Tampering (adulteração) ----------------
    Ameaca("T1", "Cadastro do lote", "T - Tampering",
           "Doador informa validade falsa para passar na validade mínima (ex.: marmita vencida)",
           "Confiar na validade digitada pelo doador", "Saúde de quem recebe",
           "Integridade / Segurança alimentar", "—", 4, 5,
           "Validade CALCULADA pelo sistema a partir de preparo + armazenamento (regras 4.2); validade mínima e rota expressa no código; questionário obrigatório com declaração assinada (regras 4.5)",
           "Inspeção da ONG na entrega (regras 6.2.1); reincidência por doador; desvio do peso/volume do próprio histórico (10.2)",
           "Suspensão do doador reincidente; a declaração assinada é a prova",
           "Implementado", 2, 4,
           "src/regras/validade.py, questionario.py; tests/test_regras.py, test_questionario.py; RECUSA_INSPECAO e BLOQUEIO_QUESTIONARIO em eventos.csv.gz"),
    Ameaca("T2", "Pipeline de treino", "T - Tampering",
           "Envenenamento de dados: lotes falsos em massa para enviesar o Modelo 2 (ex.: fazer uma região parecer 'sempre salva')",
           "Treinar com qualquer registro, sem confirmação do desfecho", "Modelos de IA, decisões de gasto",
           "Integridade", "OWASP LLM04 · MITRE ATLAS AML.T0020", 2, 4,
           "Treino só com desfechos confirmados por duas partes (ONG + transportador); dataset versionado por SHA-256; qualidade de dados no pipeline; portão de promoção com teste fixo e aprovador diferente de quem treinou",
           "Drift das entradas e saídas por PSI (10.2); checagens de qualidade de dados",
           "Rollback para a versão anterior (RegistroDeVersoes)",
           "Implementado (v2.0)", 1, 4,
           "models/*.json (sha256_dataset); src/governanca/qualidade.py; src/observabilidade/drift.py; src/models/feedback.py (portão, rollback)"),
    Ameaca("T3", "Artefatos de modelo (.joblib)", "T - Tampering",
           "Arquivo do modelo substituído por um malicioso: executa código ao ser carregado (pickle)",
           "Desserializar artefato sem verificar integridade", "Servidor da aplicação",
           "Integridade / Confidencialidade", "OWASP LLM03 · MITRE ATLAS AML.T0010", 2, 5,
           "SHA-256 do artefato gravado no treino e verificado antes do load; modelos fora do Git; imagem Docker imutável",
           "Falha de verificação gera alerta e bloqueia o start",
           "Reimplantar a partir de artefato íntegro",
           "Implementado", 1, 5,
           "src/models/artefatos.py; tests/test_artefatos.py (artefato adulterado é recusado)"),
    Ameaca("T4", "Cadeia de suprimentos (PyPI, Hugging Face, SEADE)", "T - Tampering",
           "Dependência, modelo ou fonte de dados comprometidos/alterados na origem",
           "Versões flutuantes; baixar 'a última versão' sem conferência", "Código, modelos, dados",
           "Integridade", "OWASP LLM03 · MITRE ATLAS AML.T0010", 2, 5,
           "requirements com versões fixadas; modelos HF fixados por commit, só safetensors, trust_remote_code=False, licenças inventariadas (Apache-2.0, MIT); SHA-256 fixado da fonte IPVS",
           "Trivy/Dependabot no CI para CVEs; falha de hash interrompe a curadoria",
           "Fixar versão anterior íntegra; abrir incidente",
           "Implementado (fixação e hashes); Trivy planejado CP2", 1, 5,
           "requirements.txt; src/models/nlp.py (REVISAO_EMBEDDINGS, REVISAO_NLI); src/data/curadoria_ipvs.py (SHA256_IPVS)"),
    Ameaca("T5", "Repositório / regras no código", "T - Tampering",
           "Commit malicioso ou descuidado altera pesos do ranking, tetos do caixa ou validades",
           "Push direto na main sem revisão nem testes", "Regras de negócio",
           "Integridade / Governança", "—", 2, 4,
           "Branch protegida + PR com revisão; testes de regra obrigatórios no CI",
           "Histórico do Git; teste que falha denuncia a mudança",
           "Reverter commit",
           "Parcial (testes de regra e governança); proteção de branch planejada CP2", 1, 4,
           "tests/test_regras.py"),
    Ameaca("T6", "KPIs de impacto", "T - Tampering",
           "Registros de entrega/kg inflados para maquiar o impacto (fraude de relatório social/ESG)",
           "Entrega registrada por uma só parte; tabela editável", "Relatórios de impacto, reputação",
           "Integridade", "—", 3, 3,
           "Entrega confirmada por ONG e transportador; eventos append-only; RLS no Supabase",
           "Conciliação kg doado x kg entregue por ONG",
           "Correção dos KPIs e auditoria da ONG",
           "Parcial (trilha de eventos); append-only planejado CP3", 2, 3,
           "eventos.csv.gz"),

    Ameaca("T7", "Declaração de Doação", "T - Tampering",
           "Doador ou funcionário altera a declaração depois do fato (ex.: apagar que o alimento ficou exposto em buffet)",
           "Documento editável, sem prova de integridade nem de origem", "Responsabilidade sanitária e legal",
           "Integridade / Governança", "Lei 14.016/2020", 2, 4,
           "JSON canônico + SHA-256 + assinatura RSA-PSS (3072 bits); repositório só de inclusão; correção = nova versão encadeada",
           "verificar / verificar_historico apontam conteúdo alterado ou versão quebrada",
           "A versão original continua guardada e válida",
           "Implementado (v2.0)", 1, 3,
           "src/seguranca/declaracao.py; tests/seguranca/test_declaracao.py"),
    Ameaca("T8", "Trilha de auditoria (logs)", "T - Tampering",
           "Insider apaga ou edita linhas do log para esconder desvio no caixa ou aceite indevido",
           "Log em arquivo comum, sem encadeamento; ou hash simples que o próprio insider recalcula", "Responsabilização",
           "Integridade", "MITRE ATT&CK T1070 (Indicator Removal)", 3, 4,
           "Cadeia HMAC-SHA256 com chave fora do servidor de logs; âncora de retenção também selada",
           "verificar_cadeia aponta a linha exata",
           "Restaurar do backup e investigar quem tinha acesso de escrita",
           "Implementado (v2.0)", 1, 3,
           "src/observabilidade/registro.py; tests/observabilidade/test_registro.py (alteração, remoção, chave errada, âncora forjada)"),
    Ameaca("T9", "Feedback humano (RLHF adaptado)", "T - Tampering",
           "Atacante 'corrige' a IA em massa (ex.: marmita = não perecível) para envenenar o próximo modelo",
           "Usar toda correção como rótulo de treino", "Modelos de IA",
           "Integridade", "OWASP LLM04 · MITRE ATLAS AML.T0020", 3, 4,
           "Só conta verificada; 20 correções/ator/dia; cota de 2% dos rótulos por ator; quarentena de ator com taxa de correção fora da linha de base; portão com teste fixo",
           "Quarentena gera revisão humana",
           "Descartar o lote de feedback; rollback do modelo",
           "Implementado (v2.0)", 1, 3,
           "src/models/feedback.py; tests/models/test_feedback.py; reports/feedback_demo.json (NLP real: atacante em quarentena, 0 rótulos aceitos; sem a validação o veneno passaria no portão)"),

    # ---------------- R: Repudiation (repúdio) ----------------
    Ameaca("R1", "Trilha de auditoria", "R - Repudiation",
           "ONG nega ter aceitado/recebido um lote; gestor nega ter aprovado um gasto",
           "Ações sem registro de ator e horário", "Caixa, responsabilização",
           "Governança", "—", 3, 3,
           "Todo evento registra ator (pseudônimo), ação, horário, correlação e resultado (regras 10.1); trilha de auditoria encadeada por HMAC; declaração assinada por lote",
           "verificar_cadeia aponta linha alterada, apagada ou reordenada",
           "Resolução de disputa com a trilha e a declaração como prova",
           "Implementado (v2.0)", 1, 3,
           "src/observabilidade/registro.py; tests/observabilidade/test_registro.py; eventos.csv.gz"),
    Ameaca("R2", "Decisões da IA", "R - Repudiation",
           "Não se consegue provar qual versão do modelo tomou uma decisão (ex.: acionou o pago)",
           "Modelo sem versionamento nem vínculo com os dados de treino", "Governança da IA",
           "Governança", "NIST AI RMF (Medir) · ISO/IEC 42001", 3, 3,
           "Metadados por modelo: versão, features, limiar, hash do dataset e do artefato; TODA chamada de IA registrada (modelo, versão, entradas, saída, confiança, fallback)",
           "Relatório de métricas versionado; taxa de fallback",
           "Reexecutar a decisão com a versão registrada",
           "Implementado", 1, 3,
           "models/*.json; src/models/controle.py; tests/models/test_controle.py"),

    # ---------------- I: Information Disclosure (vazamento) ----------------
    Ameaca("I1", "Dados de doador PF", "I - Information Disclosure",
           "Vazamento de CPF e endereço residencial de doadores pessoa física",
           "Documento em claro; endereço exposto antes do match; hash sem segredo (quebrável por força bruta)",
           "Dados pessoais (LGPD)", "Privacidade", "OWASP LLM02 · LGPD art. 46", 3, 5,
           "HMAC-SHA256 com chave secreta; PF no centroide do setor; endereço só após o aceite; TLS e criptografia em repouso; retenção de 12 meses",
           "Auditoria de acesso a dados pessoais",
           "Plano de resposta a incidentes + comunicação à ANPD",
           "Parcial (minimização no dataset; HMAC na regra)", 1, 4,
           "doadores.csv (sem documento em claro, PF no centroide); regras §11"),
    Ameaca("I2", "Autorização por ONG", "I - Information Disclosure",
           "Uma ONG acessa lotes, endereços ou gastos de outra ONG (IDOR)",
           "Autorização só no front-end ou feita pelo modelo", "Dados de lotes e parceiros",
           "Confidencialidade", "OWASP API1 (BOLA)", 3, 4,
           "Row Level Security no Supabase por ONG; autorização na API (fora do modelo); testes de autorização",
           "Log de acesso negado",
           "Correção da política e revisão dos acessos",
           "Planejado CP3", 1, 4, "—"),
    Ameaca("I3", "API de predição", "I - Information Disclosure",
           "Consultas repetidas para copiar o Modelo 2 ou inferir dados de treino (extração/inversão)",
           "API aberta devolvendo probabilidades brutas", "Propriedade do modelo",
           "Confidencialidade", "MITRE ATLAS AML.T0024 / AML.T0040", 2, 2,
           "Autenticação; rate limit; expor só a decisão (risco alto/baixo), não a probabilidade",
           "Detecção de padrão de consulta anômalo",
           "Bloqueio do cliente",
           "Planejado CP2", 1, 2, "Modelo treinado com dados sintéticos, sem PII (impacto baixo)"),
    Ameaca("I4", "Segredos (Supabase, tokens)", "I - Information Disclosure",
           "Chave service_role ou token vazados no repositório, nos logs ou no CI",
           "Segredo em código ou em .env versionado", "Banco inteiro",
           "Confidencialidade / Integridade", "—", 3, 5,
           ".gitignore bloqueia .env/.key/.pem; segredos no GitHub Secrets; service_role só no backend",
           "Secret scanning (gitleaks/Trivy) no CI",
           "Rotação imediata da chave vazada",
           "Parcial (.gitignore); scanning planejado CP2", 1, 5, ".gitignore"),
    Ameaca("I5", "Logs", "I - Information Disclosure",
           "Logs guardam descrições livres e dados pessoais por tempo indeterminado",
           "Registrar tudo sem minimização nem retenção (caso AtendeMais, aula 6)", "Dados pessoais",
           "Privacidade", "OWASP LLM02", 3, 3,
           "Log completo em COBERTURA e mínimo em CONTEÚDO: chaves proibidas removidas e CPF/CNPJ/e-mail/telefone mascarados no código; texto livre só como hash; retenção 90 dias (operacional) / 5 anos (auditoria)",
           "Teste que injeta dado pessoal e confere que não chega ao arquivo",
           "Expurgo por retenção (âncora selada mantém a cadeia)",
           "Implementado (v2.0)", 1, 2,
           "src/seguranca/minimizacao.py; src/observabilidade/registro.py; tests/observabilidade/test_registro.py"),
    Ameaca("I6", "Mapa / localização de ONGs", "I - Information Disclosure",
           "Exposição da localização de ONGs e de distribuições noturnas a pessoas mal-intencionadas",
           "Mapa público com pontos exatos", "Segurança física de beneficiários e voluntários",
           "Confidencialidade", "—", 2, 4,
           "Mapa público só agregado por região (heatmap); localização exata só para quem tem match aceito",
           "Revisão do que o dashboard público mostra",
           "Remover camada e revisar acessos",
           "Planejado CP3 (dashboard)", 1, 4, "regras §2.1 (índice por região)"),

    Ameaca("I7", "Mural de pedidos das ONGs", "I - Information Disclosure",
           "Mural revela padrão de necessidade e localização de ONGs e de distribuições noturnas",
           "Pedido publicado com endereço e horário", "Segurança de beneficiários e voluntários",
           "Confidencialidade", "—", 2, 3,
           "Mural agregado por região (categoria e kg), nunca com endereço; aviso só a doador que ativou, 1 por dia",
           "Revisão do que o mural expõe",
           "Remover pedido e revisar a exposição",
           "Regra definida (6.5); mural no CP3", 1, 3, "docs/regras-de-negocio.md §6.5"),

    # ---------------- D: Denial of Service (negação de serviço) ----------------
    Ameaca("D1", "Endpoint de NLP", "D - Denial of Service",
           "Entradas gigantes ou em massa esgotam a CPU e derrubam o free tier",
           "Sem limite de tamanho nem de frequência", "Disponibilidade da plataforma",
           "Disponibilidade", "OWASP LLM10 · MITRE ATLAS AML.T0029", 4, 3,
           "Guardrail limita a 500 caracteres; classificador leve (4,6 ms/texto, antes 3,3 s com NLI); limite de 120 req/min por chave + freio automático; formulário sem NLP como alternativa",
           "Linha de base de requisições e latência p95 (10.2)",
           "Desligar o NLP (o doador escolhe a categoria sozinho)",
           "Implementado (biblioteca; ligado na API no CP3)", 2, 2,
           "src/seguranca/guardrails.py; tests/seguranca/test_guardrails.py (entrada_gigante); reports/avaliacao_nlp.json (latência 3,3 s -> 4,6 ms)"),
    Ameaca("D2", "Cadastro de lotes", "D - Denial of Service",
           "Cadastro em massa de lotes falsos inunda as ONGs com ofertas (fadiga, recusas)",
           "Sem limite por doador", "Operação das ONGs",
           "Disponibilidade", "—", 3, 3,
           "Limite de requisições; doador verificado (CNPJ/CPF); peso acima de 2 t vai para revisão humana",
           "Alerta quando o volume do doador sai do PRÓPRIO histórico (mediana/MAD), antes de qualquer teto",
           "Suspender o doador e limpar as ofertas",
           "Implementado (v2.0)", 2, 2,
           "src/observabilidade/anomalias.py; tests/observabilidade/test_anomalias.py; src/regras/cadastro.py"),
    Ameaca("D3", "Infraestrutura gratuita (Render/HF)", "D - Denial of Service",
           "Serviço 'dorme' ou estoura a cota do free tier justamente num lote crítico",
           "Dependência de plano gratuito sem monitoramento", "Disponibilidade, lotes críticos",
           "Disponibilidade", "—", 3, 4,
           "Health check / keep-alive; alertas; a rota expressa cai para orientação de refrigeração",
           "Monitoramento de uptime",
           "Plano de contingência documentado",
           "Planejado CP3", 2, 3, "regras §7.5 (orientação de refrigeração como fallback)"),
    Ameaca("D4", "Caixa solidário", "D - Denial of Service",
           "Esgotamento do caixa por acionamentos pagos desnecessários ou abusivos",
           "Pagamento automático sem limites", "Recursos do caixa",
           "Disponibilidade / Governança", "OWASP LLM06", 3, 3,
           "Tetos de R$ 40 por entrega e R$ 200 por dia por ONG; gestor externo acima disso; pago só com espera estourada + risco alto",
           "Relatório de gastos por ONG",
           "Bloqueio preventivo do caixa da ONG",
           "Implementado", 1, 3,
           "src/regras/logistica.py; tests/test_regras.py (caixa, pago)"),

    Ameaca("D5", "Observabilidade / freio automático", "D - Denial of Service",
           "Atacante provoca o freio sobre um usuário legítimo (ou enche o painel de alertas para esconder um ataque real)",
           "Bloqueio automático permanente; alerta a cada evento", "Disponibilidade, atenção da equipe",
           "Disponibilidade", "—", 2, 3,
           "Freio só para sinais de segurança, reversível em 15 min e revisado por humano; sinal de negócio nunca bloqueia; alerta só quando o nível sobe na janela",
           "Freios aplicados ficam na trilha de auditoria",
           "Admin libera o freio",
           "Implementado (v2.0)", 1, 2,
           "src/observabilidade/freio.py, anomalias.py; tests/observabilidade/test_anomalias.py"),

    # ---------------- E: Elevation of Privilege (elevação de privilégio) ----------------
    Ameaca("E1", "NLP / texto livre", "E - Elevation of Privilege",
           "Prompt injection no texto do doador para manipular a classificação ou 'dar ordens' ao sistema",
           "Tratar texto do usuário como instrução", "Triagem, regras",
           "Integridade", "OWASP LLM01 · MITRE ATLAS AML.T0051", 4, 3,
           "Guardrail v2 (normalização + padrões); classificador só escolhe rótulo de lista FIXA (sem agência); doador confirma; regras revalidam",
           "Evento de bloqueio registrado para investigação",
           "Atualizar padrões a partir dos incidentes (reteste)",
           "Implementado", 3, 1,
           "reports/avaliacao_nlp.json: dataset 100% bloqueado; inéditos A 10/10 (após mitigação); B reservado 4/10 (risco residual aceito: impacto mínimo por arquitetura)"),
    Ameaca("E2", "Modelo 2 / cascata de transporte", "E - Elevation of Privilege",
           "A IA passa a executar ações de impacto (pagar entregador) sem supervisão",
           "Agência excessiva: saída do modelo vira ação", "Caixa, confiança",
           "Governança", "OWASP LLM06 (Excessive Agency)", 3, 4,
           "A IA só recomenda (agência zero); pago exige espera estourada + risco alto + tetos + confirmação da ONG ou do gestor",
           "Log de cada predição e de ACIONA_PAGO / APROVACAO_CAIXA; caixa avisa ao passar de 70% e 90% do teto",
           "Desligar o modelo por interruptor auditado (volta à heurística)",
           "Implementado", 1, 3,
           "src/regras/logistica.py; src/models/controle.py (Interruptores); docs/governanca/analise-agencia-ia.md"),
    Ameaca("E3", "Perfis e permissões (RBAC)", "E - Elevation of Privilege",
           "Usuário de ONG obtém papel de gestor/admin, ou gestor vinculado a uma ONG aprova o próprio gasto",
           "Papéis acumulados; checagem de permissão no cliente", "Caixa, cadastro de ONGs",
           "Integridade / Governança", "OWASP API5 (BFLA)", 2, 5,
           "Um perfil por usuário; segregação de funções (quem pede não aprova; admin não aprova caixa); RBAC no backend + RLS",
           "Alerta de mudança de papel",
           "Revogar o papel e revisar as aprovações",
           "Parcial (regra definida); RBAC planejado CP3", 1, 4, "regras §3 e §9"),
    Ameaca("E4", "Saída do classificador", "E - Elevation of Privilege",
           "A saída do modelo é gravada/usada por outro componente sem validação",
           "Confiar na saída da IA como dado válido", "Integridade dos lotes",
           "Integridade", "OWASP LLM05 (Improper Output Handling)", 3, 3,
           "Sugestão nunca gravada sem confirmação; saída validada contra lista fechada e [0, 1] na camada de controle; piso de segurança no M1",
           "Divergência sugestão x confirmação monitorada (vira feedback auditado)",
           "Fallback automático para a regra",
           "Implementado", 1, 2,
           "src/models/controle.py; tests/models/test_controle.py"),
    Ameaca("E6", "Validação de entrada (atribuição em massa)", "E - Elevation of Privilege",
           "Cliente envia campos extras ou tipos trocados (\"aprovado\": true, \"prioridade\": \"CRITICA\", \"peso\": \"1e9\") para pular regras",
           "Esquema permissivo que aceita campo desconhecido e converte tipos em silêncio", "Regras de negócio",
           "Integridade", "OWASP API3 (BOPLA)", 4, 4,
           "Esquemas com extra=forbid e strict; listas fechadas; faixas; datas com fuso; coerência entre campos; a prioridade nunca vem do cliente",
           "Validação rejeitada vai para o log operacional (campo e regra, sem o valor)",
           "Bloqueio pelo freio se houver rajada",
           "Implementado (v2.0)", 1, 3,
           "src/validacao/esquemas.py; tests/validacao/test_esquemas.py"),
    Ameaca("E5", "CI/CD (GitHub Actions)", "E - Elevation of Privilege",
           "Action de terceiro maliciosa ou workflow com permissão ampla publica código malicioso",
           "Actions por tag flutuante; token do workflow com escrita total", "Produção",
           "Integridade", "MITRE ATT&CK T1195 (Supply Chain Compromise)", 2, 5,
           "Actions fixadas por SHA; permissions mínimas; segredos em ambiente protegido; Bandit/Trivy como security gate",
           "Log de execuções do pipeline",
           "Revogar tokens e reimplantar",
           "Planejado CP2", 1, 4, "—"),
]

# Riscos de IA que não se encaixam em STRIDE, mas o professor cobra (aula 4: riscos éticos/operacionais).
RISCOS_IA = [
    ("A1", "Viés geográfico no ranking", "O ranking pode concentrar doações em poucas ONGs ou regiões",
     "Ética / Governança", "Peso de vulnerabilidade por setor (IPVS); monitorar kg por região no dashboard; revisão dos pesos"),
    ("A2", "Drift do Modelo 2", "Mudança de padrão (nova frota, novas ONGs) degrada as previsões",
     "Operacional", "Retreino periódico com split temporal; monitorar recall/precisão em produção; rollback"),
    ("A3", "Confiança excessiva na IA", "Equipe/ONG aceita recomendações sem questionar",
     "Operacional / Ética", "IA só recomenda; justificativa visível; humano decide ações de impacto (aula 7)"),
    ("A4", "Eficiência logística x comida perdida",
     "A preferência por pares complementares (busca x entrega) pode levar o lote a uma ONG mais distante",
     "Ética / Operacional", "Complementaridade multiplicada pelo acesso (custo medido: ~0,2 p.p. de descarte); monitorar descarte por complementaridade"),
    ("A5", "Fadiga de alertas", "Muitos alertas fazem a equipe ignorar o que importa",
     "Operacional", "Três níveis; alerta só quando o nível sobe na janela; aquecimento antes de alertar; sinais de negócio não bloqueiam"),
]


def main() -> None:
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    linhas = []
    for a in AMEACAS:
        d = asdict(a)
        d |= {"risco_inerente": a.inerente, "nivel_inerente": nivel(a.inerente),
              "risco_residual": a.residual, "nivel_residual": nivel(a.residual)}
        linhas.append(d)
    df = pd.DataFrame(linhas)
    ordem = ["id", "componente", "stride", "ameaca", "vulnerabilidade", "ativo", "principio", "referencia",
             "p", "i", "risco_inerente", "nivel_inerente", "preventivo", "detectivo", "corretivo", "status",
             "p_res", "i_res", "risco_residual", "nivel_residual", "evidencia"]
    df = df[ordem]

    DIR_SAIDA.mkdir(parents=True, exist_ok=True)
    cores = {"Crítico": "C0392B", "Alto": "E67E22", "Médio": "F1C40F", "Baixo": "27AE60"}
    with pd.ExcelWriter(DIR_SAIDA / "matriz-stride.xlsx", engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Matriz STRIDE", index=False)
        pd.DataFrame(RISCOS_IA, columns=["id", "risco", "descricao", "tipo", "controles"]).to_excel(
            xw, sheet_name="Riscos de IA", index=False)
        niveis = ["Crítico", "Alto", "Médio", "Baixo"]
        resumo = pd.DataFrame({
            "inerente": df["nivel_inerente"].value_counts().reindex(niveis, fill_value=0),
            "residual": df["nivel_residual"].value_counts().reindex(niveis, fill_value=0),
        }).rename_axis("nível")
        resumo.to_excel(xw, sheet_name="Resumo")
        ws = xw.sheets["Matriz STRIDE"]
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2C3E50")
        for col in ("L", "T"):  # nível inerente / residual
            for c in ws[col][1:]:
                c.fill, c.font = PatternFill("solid", fgColor=cores[c.value]), Font(bold=True, color="FFFFFF")
        largos = {"ameaca", "vulnerabilidade", "preventivo", "detectivo", "corretivo", "evidencia", "status"}
        for j, nome in enumerate(ordem, start=1):
            ws.column_dimensions[get_column_letter(j)].width = 42 if nome in largos else 14
        for linha in ws.iter_rows(min_row=2):
            for c in linha:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"

    def md(tabela: pd.DataFrame) -> str:
        cab = "| " + " | ".join(tabela.columns) + " |\n|" + "---|" * len(tabela.columns) + "\n"
        return cab + "\n".join("| " + " | ".join(str(v).replace("|", "/") for v in row) + " |" for row in tabela.values)

    curta = df[["id", "stride", "componente", "ameaca", "risco_inerente", "nivel_inerente",
                "preventivo", "status", "risco_residual", "nivel_residual", "evidencia"]].rename(columns={
        "risco_inerente": "P×I", "nivel_inerente": "Inerente", "risco_residual": "P×I res.", "nivel_residual": "Residual"})
    (DIR_SAIDA / "matriz-stride-tabela.md").write_text(
        "<!-- Gerado por `python -m src.seguranca.stride`. Não editar à mão. -->\n\n" + md(curta) + "\n", encoding="utf-8")
    print(resumo.to_string())


if __name__ == "__main__":
    main()
